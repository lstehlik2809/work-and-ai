"""Deliberate offline extraction; ordinary CI uses the hash-pinned JSON and Node."""
import hashlib
import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
manifest = json.loads((ROOT / 'data/ai-context-source-manifest.json').read_text('utf-8'))
source = ROOT / manifest['source']['path']
assert hashlib.sha256(source.read_bytes()).hexdigest() == manifest['source']['sha256'], 'Workbook hash changed'
workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = workbook['Table 1.12']
assert sheet.cell(1, 1).value == 'Table 1.12 Factors affecting occupational utilization, projected 2025–35'
assert [c.value for c in sheet[2]] == ['2025 National Employment Matrix occupation title', '2025 National Employment Matrix occupation code', '2025 National Employment Matrix industry title', '2025 National Employment Matrix industry code', 'Factors affecting occupational utilization']
rows = []
for index, values in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
    if values[0] == 'Source: U.S. Bureau of Labor Statistics.':
        assert index == sheet.max_row and all(v is None for v in values[1:])
        continue
    assert all(isinstance(v, str) and v.strip() for v in values), f'Unexpected row {index}'
    title, code, industry, industry_code, note = values
    rows.append(dict(row=index, title=title, code=code, industry=industry, industryCode=industry_code, note=note))
assert len(rows) == 935
result = dict(schemaVersion=1, sourceSha256=manifest['source']['sha256'], sheet='Table 1.12', startYear=2025, endYear=2035, rows=rows)
output = (json.dumps(result, ensure_ascii=False, separators=(',', ':')) + '\n').encode('utf-8')
expected = manifest['intermediate'].get('sha256')
assert expected is None or hashlib.sha256(output).hexdigest() == expected, 'Extraction differs from pinned intermediate'
(ROOT / manifest['intermediate']['path']).write_bytes(output)
print(json.dumps(dict(status='PASS', rows=len(rows), bytes=len(output), sha256=hashlib.sha256(output).hexdigest())))
