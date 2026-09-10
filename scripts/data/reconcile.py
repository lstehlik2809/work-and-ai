"""Independent raw-source oracle: does not import prepare or source-reference."""
import csv, gzip, io, json, re
from collections import defaultdict
from pathlib import Path
import openpyxl
root = Path(__file__).resolve().parents[2]
snapshot = json.loads((root/'public/data/occupations.json').read_text(encoding='utf8'))
lexicon = json.loads((root/'public/data/lexicon.json').read_text(encoding='utf8'))
def csvrows(name):
    return list(csv.DictReader(io.StringIO(gzip.decompress((root/f'data/raw/{name}.csv.gz').read_bytes()).decode('utf-8-sig'))))
book = openpyxl.load_workbook(root/'data/raw/ai-exposure-categories.xlsx',data_only=True)
sheet=book['AI Exposure Categories']
raw={str(sheet.cell(i,2).value):(i,[sheet.cell(i,c).value for c in range(1,14)]) for i in range(3,sheet.max_row+1) if sheet.cell(i,2).value}
actual={o['code']:o for o in snapshot['occupations']}
assert len(actual)==len(snapshot['occupations']) and set(actual)==set(raw)
directory=openpyxl.load_workbook(root/'data/raw/bls-directory.xlsx',data_only=True).worksheets[0]
line={directory.cell(i,2).value for i in range(5,directory.max_row+1) if directory.cell(i,4).value=='Line item'}
summary={directory.cell(i,2).value for i in range(5,directory.max_row+1) if directory.cell(i,4).value=='Summary'}
assert set(actual)==line and not set(actual)&summary
mapping=openpyxl.load_workbook(root/'data/raw/bls-crosswalk.xlsx',data_only=True)['Crosswalk']
relations={(mapping.cell(i,4).value,mapping.cell(i,2).value) for i in range(6,mapping.max_row+1)}
assert relations=={(o['code'],r['code']) for o in actual.values() for r in o['roles']}
onet={r['O*NET-SOC Code']:r for r in csvrows('occupation_data')}
tasks=defaultdict(list)
for t in csvrows('task_statements'):
    if t['Task Type']=='Core': tasks[t['O*NET-SOC Code']].append(t)
aliases={(o['title'],o['code'],None) for o in actual.values()}
for code,o in actual.items():
    row,r=raw[code]
    assert (o['sourceRow'],o['title'],o['exposure'],o['growth'],o['annualOpenings'])==(row,r[0],r[12],r[5],None if r[7] is None or r[7]=='—' else round(r[7]*1000)),code
    assert str(o['startYear']) in sheet.cell(2,3).value and str(o['endYear']) in sheet.cell(2,4).value
    for role in o['roles']:
        sr=onet[role['code']]
        assert role['title']==sr['Title'] and role['description']==sr['Description']
        selected=sorted(tasks[role['code']],key=lambda t:int(t['Task ID']))[:3]
        assert [(t['id'],t['text'],t['date']) for t in role['tasks']]==[(t['Task ID'],t['Task'],t['Date']) for t in selected]
        aliases.add((role['title'],code,role['code']))
reverse=defaultdict(set)
for c,r in relations: reverse[r].add(c)
for name,fields in [('job_titles',['Job Title','Short Title']),('sample_of_reported_titles',['Reported Job Title'])]:
    for r in csvrows(name):
        for c in reverse.get(r['O*NET-SOC Code'],()):
            for f in fields:
                if r[f].strip(): aliases.add((r[f],c,r['O*NET-SOC Code']))
assert aliases=={(a['title'],a['code'],a['onetCode']) for a in lexicon['aliases']}
exceptions=json.loads((root/'data/exceptions.json').read_text(encoding='utf8'))
assert set(onet)-{r for c,r in relations}=={r['code'] for r in exceptions['unmappedOnet']}
assert set(exceptions['excludedSummaryCodes'])==summary
assert ('11-1011','11-1011.03') in relations
assert next(r for r in actual['11-1011']['roles'] if r['code']=='11-1011.03')['broader'] is True
# Twelve fixed source rows plus any categories/signs not represented in that set.
codes=['11-1011','11-2011','13-2011','15-1252','17-2141','19-3032','29-1141','31-9091','41-2011','43-4161','47-2111','53-3032']
for cat in ['Low','Moderate','High','Very high']:
    if not any(actual[c]['exposure']==cat for c in codes): codes.append(next(c for c,o in actual.items() if o['exposure']==cat))
for sign in [-1,0,1]:
    if not any((actual[c]['growth']>0)-(actual[c]['growth']<0)==sign for c in codes): codes.append(next(c for c,o in actual.items() if o['growth'] is not None and (o['growth']>0)-(o['growth']<0)==sign))
report={'status':'PASS','oracle':'independent direct openpyxl/csv reads; no prepare imports or generated reference','fullPopulation':len(actual),'relations':len(relations),'aliases':len(aliases),'excludedSummaries':len(summary),'unmappedOnet':len(exceptions['unmappedOnet']),'sourceSamples':[{'sheet':sheet.title,**{k:actual[c][k] for k in ['code','title','sourceRow','exposure','growth','annualOpenings','startYear','endYear']}} for c in codes]}
(root/'data/reconciliation-report.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf8')
print(json.dumps(report))
